"""
Excel导出路由
"""

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import Optional
import io
from datetime import datetime
from urllib.parse import quote
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from backend.database import get_db
from backend.utils.response import success_response
from backend.utils.exceptions import ResourceNotFoundException
from backend.models import Task, Evaluation, Member, Week, Achievement, WeeklyReport
from backend.config import TASK_STATUS, TASK_TYPES, TASK_DIFFICULTY, EVALUATION_LEVELS
from backend.services.llm_service import LLMService
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

router = APIRouter()

def create_excel_response(filename: str, workbook: openpyxl.Workbook):
    """创建Excel文件响应"""
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    # 对中文文件名进行URL编码
    encoded_filename = quote(filename)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"}
    )

def style_header(ws, headers: list):
    """设置表头样式"""
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

@router.get("/export/week-plan", summary="导出周计划表")
def export_week_plan(
    week_id: int = Query(..., description="周次ID"),
    db: Session = Depends(get_db)
):
    """导出周计划表"""
    week = db.query(Week).filter(Week.id == week_id).first()
    if not week:
        raise ResourceNotFoundException("周次", week_id)

    tasks = db.query(Task).filter(Task.week_id == week_id).order_by(Task.deadline).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "周计划"

    headers = ["序号", "任务名称", "任务类型", "难度", "责任人", "截止时间", "交付要求", "状态", "是否延期"]
    style_header(ws, headers)

    for idx, task in enumerate(tasks, 1):
        row = [
            idx,
            task.name,
            TASK_TYPES.get(task.task_type, task.task_type),
            TASK_DIFFICULTY.get(task.difficulty, task.difficulty),
            task.assignee.name if task.assignee else "",
            task.deadline.strftime("%Y-%m-%d") if task.deadline else "",
            task.delivery_requirement or "",
            TASK_STATUS.get(task.status, task.status),
            "是" if task.is_overdue else "否"
        ]
        for col, value in enumerate(row, 1):
            ws.cell(row=idx + 1, column=col, value=value)

    # 调整列宽
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15

    filename = f"周计划_{week.name}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return create_excel_response(filename, wb)

@router.get("/export/week-evaluation", summary="导出周评价汇总")
def export_week_evaluation(
    week_id: int = Query(..., description="周次ID"),
    db: Session = Depends(get_db)
):
    """导出周评价汇总"""
    week = db.query(Week).filter(Week.id == week_id).first()
    if not week:
        raise ResourceNotFoundException("周次", week_id)

    tasks = db.query(Task).filter(Task.week_id == week_id).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "周评价汇总"

    headers = ["序号", "任务名称", "责任人", "评价等级", "评价意见", "基础积分", "加分", "最终积分"]
    style_header(ws, headers)

    row_num = 2
    idx = 1
    for task in tasks:
        if task.evaluations:
            for eval in task.evaluations:
                ws.cell(row=row_num, column=1, value=idx)
                ws.cell(row=row_num, column=2, value=task.name)
                ws.cell(row=row_num, column=3, value=task.assignee.name if task.assignee else "")
                ws.cell(row=row_num, column=4, value=EVALUATION_LEVELS.get(eval.level, eval.level))
                ws.cell(row=row_num, column=5, value=eval.comment or "")
                ws.cell(row=row_num, column=6, value=eval.base_score)
                ws.cell(row=row_num, column=7, value=eval.bonus_score)
                ws.cell(row=row_num, column=8, value=eval.final_score)
                row_num += 1
                idx += 1
        else:
            ws.cell(row=row_num, column=1, value=idx)
            ws.cell(row=row_num, column=2, value=task.name)
            ws.cell(row=row_num, column=3, value=task.assignee.name if task.assignee else "")
            ws.cell(row=row_num, column=4, value="未评价")
            ws.cell(row=row_num, column=5, value="")
            ws.cell(row=row_num, column=6, value=0)
            ws.cell(row=row_num, column=7, value=0)
            ws.cell(row=row_num, column=8, value=0)
            row_num += 1
            idx += 1

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15

    filename = f"周评价汇总_{week.name}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return create_excel_response(filename, wb)

@router.get("/export/member-statistics", summary="导出个人贡献统计")
def export_member_statistics(
    week_id: Optional[int] = Query(None, description="周次ID"),
    db: Session = Depends(get_db)
):
    """导出个人贡献统计"""
    if week_id:
        week = db.query(Week).filter(Week.id == week_id).first()
        week_name = week.name if week else ""
    else:
        week_name = "全部"

    members = db.query(Member).filter(Member.status == "active").all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "个人贡献统计"

    headers = ["姓名", "所属单位", "本周任务数", "已完成数", "延期数", "优秀数", "本周积分", "累计积分"]
    style_header(ws, headers)

    for idx, member in enumerate(members, 2):
        # 本周统计
        query = db.query(Task).filter(Task.assignee_id == member.id)
        if week_id:
            query = query.filter(Task.week_id == week_id)
        tasks = query.all()

        weekly_total = len(tasks)
        weekly_completed = len([t for t in tasks if t.status == "completed"])
        weekly_overdue = len([t for t in tasks if t.is_overdue])
        weekly_excellent = 0
        weekly_score = 0

        for task in tasks:
            for eval in task.evaluations:
                weekly_score += eval.final_score
                if eval.level == "excellent":
                    weekly_excellent += 1

        # 累计积分
        total_score = sum([
            sum([e.final_score for e in t.evaluations])
            for t in db.query(Task).filter(Task.assignee_id == member.id).all()
        ])

        ws.cell(row=idx, column=1, value=member.name)
        ws.cell(row=idx, column=2, value=member.unit)
        ws.cell(row=idx, column=3, value=weekly_total)
        ws.cell(row=idx, column=4, value=weekly_completed)
        ws.cell(row=idx, column=5, value=weekly_overdue)
        ws.cell(row=idx, column=6, value=weekly_excellent)
        ws.cell(row=idx, column=7, value=weekly_score)
        ws.cell(row=idx, column=8, value=total_score)

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15

    filename = f"个人贡献统计_{week_name}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return create_excel_response(filename, wb)

@router.get("/export/achievements", summary="导出成果清单")
def export_achievements(
    week_id: Optional[int] = Query(None, description="周次ID"),
    keyword: Optional[str] = Query(None, description="关键词搜索"),
    achievement_type: Optional[str] = Query(None, description="成果类型筛选"),
    member_id: Optional[int] = Query(None, description="提交人筛选"),
    is_excellent: Optional[bool] = Query(None, description="优秀成果筛选"),
    db: Session = Depends(get_db)
):
    """导出成果清单（从成果库读取，支持当前筛选条件）"""
    query = db.query(Achievement)

    if week_id:
        query = query.filter(Achievement.week_id == week_id)
    if keyword:
        query = query.filter(Achievement.name.like(f"%{keyword}%"))
    if achievement_type:
        query = query.filter(Achievement.achievement_type == achievement_type)
    if member_id:
        query = query.filter(Achievement.member_id == member_id)
    if is_excellent is not None:
        query = query.filter(Achievement.is_excellent == is_excellent)

    achievements = query.order_by(Achievement.created_at.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "成果清单"

    headers = ["序号", "成果名称", "成果说明", "成果链接", "成果类型", "提交人", "所属周次", "提交时间", "是否优秀"]
    style_header(ws, headers)

    for idx, a in enumerate(achievements, 2):
        ws.cell(row=idx, column=1, value=idx - 1)
        ws.cell(row=idx, column=2, value=a.name)
        ws.cell(row=idx, column=3, value=a.description or "")
        ws.cell(row=idx, column=4, value=a.link or "")
        ws.cell(row=idx, column=5, value=a.achievement_type or "")
        ws.cell(row=idx, column=6, value=a.member.name if a.member else "")
        ws.cell(row=idx, column=7, value=a.week.name if a.week else "")
        ws.cell(row=idx, column=8, value=a.created_at.strftime("%Y-%m-%d %H:%M") if a.created_at else "")
        ws.cell(row=idx, column=9, value="★" if a.is_excellent else "")

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18

    filename = f"成果清单_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return create_excel_response(filename, wb)


@router.get("/export/weekly-report-word", summary="导出专班周报 Word")
def export_weekly_report_word(
    week_id: int = Query(..., description="周次ID"),
    db: Session = Depends(get_db)
):
    """导出专班周报 Word 文档"""
    week = db.query(Week).filter(Week.id == week_id).first()
    if not week:
        raise ResourceNotFoundException("周次", week_id)

    reports = db.query(WeeklyReport).filter(WeeklyReport.week_id == week_id).all()

    items = []
    for r in reports:
        item = {
            "member_name": r.member.name if r.member else None,
            "work_content": r.work_content,
            "main_results": r.main_results,
            "problems": r.problems,
            "next_week_plan": r.next_week_plan,
        }
        if r.analysis:
            item["analysis"] = {
                "contribution_summary": r.analysis.contribution_summary,
                "risk_alerts": r.analysis.risk_alerts,
            }
        items.append(item)

    # 调用大模型生成专班周报草稿
    llm_service = LLMService()
    summary = llm_service.generate_team_summary(items, week.name)

    # 构建 Word 文档
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "SimSun"
    style.font.size = Pt(11)
    style.paragraph_format.line_spacing = 1.5

    # 标题
    title = doc.add_heading(f"专班周报 — {week.name}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 日期
    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date_para.add_run(f"生成日期：{datetime.now().strftime('%Y-%m-%d')}")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(128, 128, 128)

    doc.add_paragraph()  # 空行

    # AI 生成摘要
    if summary.get("summary_text"):
        doc.add_heading("一、本周总体情况", level=1)
        for line in summary["summary_text"].strip().split("\n"):
            line = line.strip()
            if not line:
                continue
            if line.startswith("**") and line.endswith("**"):
                doc.add_heading(line.strip("**"), level=2)
            elif line.startswith("#"):
                doc.add_heading(line.lstrip("#").strip(), level=2)
            else:
                doc.add_paragraph(line)

        doc.add_paragraph()

    # 成员详情
    doc.add_heading("二、成员周报详情", level=1)
    for i, item in enumerate(items, 1):
        doc.add_heading(f"{i}. {item['member_name'] or '未知'}", level=2)

        if item.get("work_content"):
            p = doc.add_paragraph()
            run = p.add_run("本周工作：")
            run.bold = True
            p.add_run(item["work_content"])

        if item.get("main_results"):
            p = doc.add_paragraph()
            run = p.add_run("主要成果：")
            run.bold = True
            p.add_run(item["main_results"])

        if item.get("problems"):
            p = doc.add_paragraph()
            run = p.add_run("存在问题：")
            run.bold = True
            p.add_run(item["problems"])

        if item.get("next_week_plan"):
            p = doc.add_paragraph()
            run = p.add_run("下周计划：")
            run.bold = True
            p.add_run(item["next_week_plan"])

        analysis = item.get("analysis") or {}
        if analysis.get("contribution_summary"):
            p = doc.add_paragraph()
            run = p.add_run("贡献摘要：")
            run.bold = True
            p.add_run(analysis["contribution_summary"])

        if analysis.get("risk_alerts"):
            p = doc.add_paragraph()
            run = p.add_run("风险提示：")
            run.bold = True
            for alert in analysis["risk_alerts"]:
                doc.add_paragraph(alert, style="List Bullet")

        doc.add_paragraph()  # 成员间分隔

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)

    filename = quote(f"专班周报_{week.name}_{datetime.now().strftime('%Y%m%d')}.docx")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"}
    )
