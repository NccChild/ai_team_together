"""Excel导出服务"""

from sqlalchemy.orm import Session
from typing import Optional
from datetime import datetime
import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from fastapi.responses import StreamingResponse
from backend.models import Task, Evaluation, Member, Delivery, Week
from backend.utils.exceptions import ResourceNotFoundException
from backend.config import TASK_STATUS, TASK_TYPES, TASK_DIFFICULTY, EVALUATION_LEVELS


def _create_excel_response(filename: str, workbook: openpyxl.Workbook):
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


def _style_header(ws, headers: list):
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")


def export_week_plan(db: Session, week_id: int):
    week = db.query(Week).filter(Week.id == week_id).first()
    if not week:
        raise ResourceNotFoundException("周次", week_id)

    tasks = db.query(Task).filter(Task.week_id == week_id).order_by(Task.deadline).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "周计划"

    headers = ["序号", "任务名称", "任务类型", "难度", "责任人", "截止时间", "交付要求", "状态", "是否延期"]
    _style_header(ws, headers)

    for idx, task in enumerate(tasks, 1):
        row = [
            idx,
            task.name,
            TASK_TYPES.get(task.task_type, task.task_type),
            TASK_DIFFICULTY.get(task.difficulty, task.difficulty),
            task.assignee.name if task.assignee else "",
            task.deadline.strftime("%Y-%m-%d") if task.deadline else "",
            task.delivery_requirement or "",
            TASK_STATUS.get(task.status, task.status),
            "是" if task.is_overdue else "否"
        ]
        for col, value in enumerate(row, 1):
            ws.cell(row=idx + 1, column=col, value=value)

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15

    filename = f"周计划_{week.name}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return _create_excel_response(filename, wb)


def export_week_evaluation(db: Session, week_id: int):
    week = db.query(Week).filter(Week.id == week_id).first()
    if not week:
        raise ResourceNotFoundException("周次", week_id)

    tasks = db.query(Task).filter(Task.week_id == week_id).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "周评价汇总"

    headers = ["序号", "任务名称", "责任人", "评价等级", "评价意见", "基础积分", "加分", "最终积分"]
    _style_header(ws, headers)

    row_num = 2
    idx = 1
    for task in tasks:
        if task.evaluations:
            for eval in task.evaluations:
                ws.cell(row=row_num, column=1, value=idx)
                ws.cell(row=row_num, column=2, value=task.name)
                ws.cell(row=row_num, column=3, value=task.assignee.name if task.assignee else "")
                ws.cell(row=row_num, column=4, value=EVALUATION_LEVELS.get(eval.level, eval.level))
                ws.cell(row=row_num, column=5, value=eval.comment or "")
                ws.cell(row=row_num, column=6, value=eval.base_score)
                ws.cell(row=row_num, column=7, value=eval.bonus_score)
                ws.cell(row=row_num, column=8, value=eval.final_score)
                row_num += 1
                idx += 1
        else:
            ws.cell(row=row_num, column=1, value=idx)
            ws.cell(row=row_num, column=2, value=task.name)
            ws.cell(row=row_num, column=3, value=task.assignee.name if task.assignee else "")
            ws.cell(row=row_num, column=4, value="未评价")
            ws.cell(row=row_num, column=5, value="")
            ws.cell(row=row_num, column=6, value=0)
            ws.cell(row=row_num, column=7, value=0)
            ws.cell(row=row_num, column=8, value=0)
            row_num += 1
            idx += 1

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15

    filename = f"周评价汇总_{week.name}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return _create_excel_response(filename, wb)


def export_member_statistics(db: Session, week_id: Optional[int] = None):
    if week_id:
        week = db.query(Week).filter(Week.id == week_id).first()
        week_name = week.name if week else ""
    else:
        week_name = "全部"

    members = db.query(Member).filter(Member.status == "active").all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "个人贡献统计"

    headers = ["姓名", "所属单位", "本周任务数", "已完成数", "延期数", "优秀数", "本周积分", "累计积分"]
    _style_header(ws, headers)

    for idx, member in enumerate(members, 2):
        query = db.query(Task).filter(Task.assignee_id == member.id)
        if week_id:
            query = query.filter(Task.week_id == week_id)
        tasks = query.all()

        weekly_total = len(tasks)
        weekly_completed = len([t for t in tasks if t.status == "completed"])
        weekly_overdue = len([t for t in tasks if t.is_overdue])
        weekly_excellent = 0
        weekly_score = 0

        for task in tasks:
            for eval in task.evaluations:
                weekly_score += eval.final_score
                if eval.level == "excellent":
                    weekly_excellent += 1

        total_score = sum([
            sum([e.final_score for e in t.evaluations])
            for t in db.query(Task).filter(Task.assignee_id == member.id).all()
        ])

        ws.cell(row=idx, column=1, value=member.name)
        ws.cell(row=idx, column=2, value=member.unit)
        ws.cell(row=idx, column=3, value=weekly_total)
        ws.cell(row=idx, column=4, value=weekly_completed)
        ws.cell(row=idx, column=5, value=weekly_overdue)
        ws.cell(row=idx, column=6, value=weekly_excellent)
        ws.cell(row=idx, column=7, value=weekly_score)
        ws.cell(row=idx, column=8, value=total_score)

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15

    filename = f"个人贡献统计_{week_name}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return _create_excel_response(filename, wb)


def export_achievements(db: Session, week_id: Optional[int] = None):
    query = db.query(Delivery).filter(Delivery.is_latest == True)
    if week_id:
        query = query.join(Task, Delivery.task_id == Task.id).filter(Task.week_id == week_id)

    deliveries = query.all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "成果清单"

    headers = ["序号", "成果名称", "成果说明", "成果链接", "成果类型", "提交人", "提交时间", "是否优秀"]
    _style_header(ws, headers)

    for idx, d in enumerate(deliveries, 2):
        ws.cell(row=idx, column=1, value=idx - 1)
        ws.cell(row=idx, column=2, value=d.name)
        ws.cell(row=idx, column=3, value=d.description or "")
        ws.cell(row=idx, column=4, value=d.link or "")
        ws.cell(row=idx, column=5, value=d.delivery_type or "")
        ws.cell(row=idx, column=6, value=d.submitter.name if d.submitter else "")
        ws.cell(row=idx, column=7, value=d.submitted_at.strftime("%Y-%m-%d %H:%M") if d.submitted_at else "")
        ws.cell(row=idx, column=8, value="")

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18

    filename = f"成果清单_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return _create_excel_response(filename, wb)
